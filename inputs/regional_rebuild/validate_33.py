# -*- coding: utf-8 -*-
"""
R1e step 1-2: rebuild the 47 prefecture composition profiles from the raw MoEJ FY2021 facility
table with the rules R1-R4 of RULES_prefecture_composition.md (written before this run), and
compare them item by item with the old compiled workbook (V1-V3).  Reads only; writes only
into this folder:  prefecture_rebuilt_raw.csv, validate_33_report.csv, validate_33.log (via tee).
Nothing in the calculation chain is touched.
"""
import csv
import hashlib
import os
import statistics as st
import sys
import warnings

import openpyxl

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(encoding='utf-8')
HERE = os.path.dirname(os.path.abspath(__file__))
R1 = os.path.normpath(os.path.join(HERE, '..', '..'))
FAC = os.path.join(R1, '03_new_sources', 'MoEJ_FY2021', 'r3_seibi_facility_01.xlsx')
OLD = os.path.normpath(os.path.join(R1, '..', 'regions', 'source_workbook', '全国ごみ性状分析　導入効果把握.xlsx'))
assert hashlib.sha256(open(FAC, 'rb').read()).hexdigest().startswith('77a95fe0'), 'raw facility file changed'

KEYS = ['paper', 'plas', 'wood', 'kit', 'inert', 'other']
TOL = 1e-6

# ---------------------------------------------------------------- raw facility table
ws = openpyxl.load_workbook(FAC, read_only=True, data_only=True)['焼却']
rows = list(ws.iter_rows(min_row=1, max_row=1066, values_only=True))
COL = {'pref': 0, 'thr': 6, 'tot': 55, 'paper': 56, 'plas': 57, 'wood': 58, 'kit': 59, 'inert': 60,
       'other': 61, 'water': 64}
HDR = {'pref': '都道府県名', 'thr': '年間処理量', 'tot': '合計', 'paper': '紙・布類',
       'plas': 'ﾋﾞﾆｰﾙ、合成樹脂、ｺﾞﾑ、皮革類', 'wood': '木、竹、わら類', 'kit': 'ちゅう芥類',
       'inert': '不燃物類', 'other': 'その他', 'water': '水分'}
for k, h in HDR.items():
    seen = [str(rows[r][COL[k]]) for r in range(1, 5)]
    assert h in seen, (k, h, seen)
fac = [r for r in rows[6:] if r[0]]
assert len(fac) == 1060

num = lambda v: float(v) if isinstance(v, (int, float)) else None
by = {}
n_r3, n_water_missing, n_in = 0, 0, 0
for r in fac:
    tot = num(r[COL['tot']])
    if not (tot and tot > 0):                                   # R1
        continue
    n_in += 1
    d = {k: (num(r[COL[k]]) or 0.0) for k in KEYS}               # R2
    five = sum(d[k] for k in KEYS[:5])
    if abs(five + d['other'] - 100) > 1.5:                       # R3
        d['other'] = max(0.0, 100.0 - five)
        n_r3 += 1
    w = num(r[COL['water']])
    if w is None:
        n_water_missing += 1
    d['water'] = w
    by.setdefault(r[0], []).append(d)
assert len(by) == 47, len(by)

reb = {}
for p, L in by.items():                                          # R4
    m = {k: st.mean(x[k] for x in L) for k in KEYS}
    ws_ = [x['water'] for x in L if x['water'] is not None]
    m['moisture'] = st.mean(ws_) if ws_ else None
    m['n_fac'] = len(L)
    m['n_water'] = len(ws_)
    reb[p] = m

# ---------------------------------------------------------------- old compiled workbook
wo = openpyxl.load_workbook(OLD, read_only=True, data_only=True)['日本全国ごみ性状']
old = {}
for r in wo.iter_rows(min_row=1, max_row=80, values_only=True):
    if not r[1] or not isinstance(r[3], (int, float)):
        continue
    old[str(r[1]).strip()] = dict(paper=r[3], plas=r[4], wood=r[5], kit=r[6], inert=r[7], other=r[8],
                                  moisture=r[11])
assert len(old) == 47 and set(old) == set(reb), set(old) ^ set(reb)

# ---------------------------------------------------------------- V1 / V2
ok, bad, rep = [], [], []
for p in old:
    diffs = {k: reb[p][k] - old[p][k] for k in KEYS + ['moisture']}
    mx = max(abs(v) for v in diffs.values())
    (ok if mx <= TOL else bad).append(p)
    rep.append([p, reb[p]['n_fac'], reb[p]['n_water'], 'yes' if mx <= TOL else 'no']
               + [round(diffs[k], 6) for k in KEYS + ['moisture']] + [round(mx, 6)])
print(f'V3  facilities included (BD > 0): {n_in} of {len(fac)}; R3 triggered {n_r3} times; '
      f'moisture blank in {n_water_missing} included facilities')
print(f'V1  prefectures reproduced item by item (|diff| <= {TOL:g} on six fractions and moisture): '
      f'{len(ok)} of 47')
print('    reproduced:     ' + ' '.join(sorted(ok)))
print('    not reproduced: ' + ' '.join(sorted(bad)))
print('V2  differences of the non-reproduced prefectures (rebuilt - old, pp):')
print('    ' + f'{"prefecture":<8}{"n":>4}' + ''.join(f'{k:>9}' for k in KEYS + ['moisture']))
for row in sorted(rep, key=lambda x: -x[-1]):
    if row[3] == 'no':
        print('    ' + f'{row[0]:<8}{row[1]:>4}' + ''.join(f'{v:>+9.2f}' for v in row[4:11]))

with open(os.path.join(HERE, 'validate_33_report.csv'), 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['prefecture', 'n_facilities_BD_gt_0', 'n_with_moisture', 'reproduced']
               + [f'diff_{k}_pp' for k in KEYS + ['moisture']] + ['max_abs_diff'])
    w.writerows(sorted(rep))
with open(os.path.join(HERE, 'prefecture_rebuilt_raw.csv'), 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['prefecture', 'n_facilities', 'n_with_moisture'] + [f'{k}_pct_raw_mean' for k in KEYS]
               + ['moisture_pct'])
    for p in sorted(reb):
        w.writerow([p, reb[p]['n_fac'], reb[p]['n_water']] + [reb[p][k] for k in KEYS] + [reb[p]['moisture']])
print('wrote validate_33_report.csv, prefecture_rebuilt_raw.csv (raw means, before R5 normalisation)')
