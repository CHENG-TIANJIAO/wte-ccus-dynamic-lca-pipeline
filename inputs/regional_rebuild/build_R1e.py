# -*- coding: utf-8 -*-
"""
R1e step 3: transparent rebuild of the regional layer from the raw MoEJ FY2021 files, with the
rules of RULES_prefecture_composition.md (R1-R7, R4 as revised 2026-09-20).  Writes ONLY into
this folder; the calculation chain is not touched.

Outputs (this folder)
  prefecture_waste_carbon_R1e.csv   47 prefectures, same columns as the chain's
                                    regional_source/prefecture_waste_carbon.csv (+ n_facilities,
                                    n_moisture, direct_incineration_FY2021_t)
  region_waste_carbon_R1e_W2a.csv   BASE: 10 regions, prefecture values weighted by 直接焼却量;
                                    same columns as the chain's region_waste_carbon.csv
  region_waste_carbon_R1e_W1.csv    sensitivity: facility level (BD > 0 and 年間処理量 > 0),
                                    weighted by facility 年間処理量
  (sensitivity U = the chain's region_waste_carbon_unweighted.csv, the submitted basis; not rebuilt)
  R1e_gsd.json                      the two MC widths, exp(sample sd of ln x, ddof = 1), full precision
  R1e_comparison.csv / .md          old (R1d W2a, submitted U) vs new, per region; GSDs; preview of
                                    the 2025 regional results through the chain's own regional_model
Run:  python build_R1e.py
"""
import csv
import hashlib
import io
import json
import os
import sys
import warnings

import numpy as np
import openpyxl

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(encoding='utf-8')
HERE = os.path.dirname(os.path.abspath(__file__))
ANA = os.path.dirname(HERE)
R1 = os.path.dirname(ANA)
CODE = os.path.join(ANA, 'chain_r1', 'calculation_code')
# the R1d state of the chain (old values for the comparison): the backup taken before R1e was installed
PRE = os.path.join(ANA, 'chain_r1', '_snapshots', 'pre_R1e_2026-09-20', 'calculation_code')
MOEJ = os.path.join(R1, '03_new_sources', 'MoEJ_FY2021')
FAC = os.path.join(MOEJ, 'r3_seibi_facility_01.xlsx')
TOT = os.path.join(MOEJ, 'r3_shori_total_02.xlsx')
for f, h in ((FAC, '77a95fe0dbe53dc1'), (TOT, 'b12aefdffcf1ec30')):
    assert hashlib.sha256(open(f, 'rb').read()).hexdigest().startswith(h), f

# the chain's own constants and carbon function (read-only import; build_regions needs a
# source path at import time, the old compiled workbook is given only to satisfy that)
os.environ['CCUS_SRC_WASTE'] = os.path.join(os.path.dirname(R1), 'regions', 'source_workbook',
                                            '全国ごみ性状分析　導入効果把握.xlsx')
sys.path.insert(0, os.path.join(CODE, 'regional_source'))
sys.argv = [sys.argv[0]]
import build_regions as BR                                     # noqa: E402

KEYS = ['paper', 'plas', 'wood', 'kit', 'inert', 'other']
num = lambda v: float(v) if isinstance(v, (int, float)) else None
gsd = lambda x: float(np.exp(np.std(np.log(np.asarray(x, float)), ddof=1)))


def normalise(d):
    """R5 (identical to build_regions.load()): residual guard, then normalise to 100 %."""
    five = sum(d[k] for k in KEYS[:5])
    d['residual_fixed'] = False
    if abs(five + d['other'] - 100) > 1.5:
        d['other'] = max(0.0, 100.0 - five)
        d['residual_fixed'] = True
    tot = five + d['other']
    for k in KEYS:
        d[k] = d[k] * 100.0 / tot
    return d


def with_carbon(d):
    """R6 (identical to build_regions.carbon() and main())."""
    totC, fosC = BR.carbon(d)
    d['totC'], d['fosC'] = totC, fosC
    d['fossil_share'] = 100 * fosC / totC
    d['CO2_total'] = totC / 100 * 44 / 12
    d['CO2_fossil'] = fosC / 100 * 44 / 12
    return d


# ---------------------------------------------------------------- raw tables
ws = openpyxl.load_workbook(FAC, read_only=True, data_only=True)['焼却']
rows = list(ws.iter_rows(min_row=1, max_row=1066, values_only=True))
COL = {'pref': 0, 'name': 4, 'thr': 6, 'tot': 55, 'paper': 56, 'plas': 57, 'wood': 58, 'kit': 59,
       'inert': 60, 'other': 61, 'water': 64}
HDR = {'pref': '都道府県名', 'name': '施設名称', 'thr': '年間処理量', 'tot': '合計', 'paper': '紙・布類',
       'plas': 'ﾋﾞﾆｰﾙ、合成樹脂、ｺﾞﾑ、皮革類', 'wood': '木、竹、わら類', 'kit': 'ちゅう芥類',
       'inert': '不燃物類', 'other': 'その他', 'water': '水分'}
for k, h in HDR.items():
    assert h in [str(rows[r][COL[k]]) for r in range(1, 5)], (k, h)
assert rows[5][COL['thr']] == '（t/年度）'
fac_all = [r for r in rows[6:] if r[0]]
assert len(fac_all) == 1060

wt = openpyxl.load_workbook(TOT, read_only=True, data_only=True)['ごみ処理概要']
assert wt['P3'].value == '直接焼却量' and wt['P6'].value == '（ｔ）'
direct = {wt.cell(r, 1).value: float(wt.cell(r, 16).value) for r in range(7, 54)}
NAT = float(wt.cell(54, 16).value)
assert len(direct) == 47 and abs(sum(direct.values()) - NAT) < 0.5 and abs(NAT - 31491040) < 0.5
assert set(direct) == set(BR.REGION)

# ---------------------------------------------------------------- facilities (R1-R3)
facs, n_r3 = [], 0
for r in fac_all:
    t = num(r[COL['tot']])
    if not (t and t > 0):                                                      # R1
        continue
    d = {k: (num(r[COL[k]]) or 0.0) for k in KEYS}                              # R2
    five = sum(d[k] for k in KEYS[:5])
    if abs(five + d['other'] - 100) > 1.5:                                      # R3
        d['other'] = max(0.0, 100.0 - five)
        n_r3 += 1
    w = num(r[COL['water']])
    d['water'] = w if (w is not None and w > 0) else None                       # R4 (revised)
    d['thr'] = num(r[COL['thr']]) or 0.0
    d['pref'] = r[0]
    facs.append(d)
assert n_r3 == 0

# ---------------------------------------------------------------- prefectures (R4-R6)
pref = {}
for p in BR.REGION:
    L = [f for f in facs if f['pref'] == p]
    assert L, p
    d = {k: float(np.mean([f[k] for f in L])) for k in KEYS}
    wv = [f['water'] for f in L if f['water'] is not None]
    d['moisture'] = float(np.mean(wv)) if wv else None
    d['n_fac'], d['n_moist'] = len(L), len(wv)
    d['pref'], d['region'] = p, BR.REGION[p]
    pref[p] = with_carbon(normalise(d))
missing_moist = [p for p in pref if pref[p]['moisture'] is None]

# ---------------------------------------------------------------- regions
OLD_W2A = {r['region']: r for r in csv.DictReader(open(os.path.join(PRE, 'region_waste_carbon.csv'),
                                                      encoding='utf-8-sig'))}
OLD_U = {r['region']: r for r in csv.DictReader(open(os.path.join(PRE, 'region_waste_carbon_unweighted.csv'),
                                                    encoding='utf-8-sig'))}
REGIONS = sorted(OLD_W2A)
HEAD = ['region', 'n_prefectures', 'plastic_pct', 'moisture_pct', 'fossil_share_of_C_pct',
        'CO2_total_t_per_t_dry', 'CO2_fossil_t_per_t_dry', 'grid_EF_2025_tCO2_per_kWh',
        'grid_EF_2050_proportional', 'grid_EF_2050_convergence', 'direct_incineration_FY2021_t',
        'share_of_national_direct_incineration_pct']


def wmean(vals, w):
    vals, w = np.asarray(vals, float), np.asarray(w, float)
    return float((vals * w).sum() / w.sum())


def region_rows(kind):
    out = {}
    for reg in REGIONS:
        ms = [pref[p] for p in pref if pref[p]['region'] == reg]
        D = sum(direct[m['pref']] for m in ms)
        if kind == 'W2a':
            w = [direct[m['pref']] for m in ms]
            v = {k: wmean([m[k] for m in ms], w) for k in ('plas', 'fossil_share', 'CO2_total', 'CO2_fossil')}
            mw = [(m['moisture'], wi) for m, wi in zip(ms, w) if m['moisture'] is not None]
            v['moisture'] = wmean([a for a, _ in mw], [b for _, b in mw])
        else:                                                                     # W1
            L = [with_carbon(normalise(dict(f))) for f in facs if BR.REGION[f['pref']] == reg and f['thr'] > 0]
            w = [f['thr'] for f in L]
            v = {k: wmean([f[k] for f in L], w) for k in ('plas', 'fossil_share', 'CO2_total', 'CO2_fossil')}
            mw = [(f['water'], f['thr']) for f in L if f['water'] is not None]
            v['moisture'] = wmean([a for a, _ in mw], [b for _, b in mw])
            v['n_fac'] = len(L)
        o = OLD_W2A[reg]
        out[reg] = dict(v, n=len(ms), D=D, ef=[o['grid_EF_2025_tCO2_per_kWh'], o['grid_EF_2050_proportional'],
                                               o['grid_EF_2050_convergence']])
    return out


def write_region(fn, R):
    with open(os.path.join(HERE, fn), 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(HEAD)
        for reg in REGIONS:
            r = R[reg]
            w.writerow([reg, r['n'], round(r['plas'], 2), round(r['moisture'], 2), round(r['fossil_share'], 2),
                        round(r['CO2_total'], 4), round(r['CO2_fossil'], 4)] + r['ef']
                       + [int(round(r['D'])), round(100 * r['D'] / NAT, 2)])


W2A, W1 = region_rows('W2a'), region_rows('W1')
write_region('region_waste_carbon_R1e_W2a.csv', W2A)
write_region('region_waste_carbon_R1e_W1.csv', W1)
n_w1 = sum(W1[r]['n_fac'] for r in REGIONS)
thr_w1 = sum(f['thr'] for f in facs if f['thr'] > 0)

with open(os.path.join(HERE, 'prefecture_waste_carbon_R1e.csv'), 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['prefecture', 'region', 'paper_pct', 'plastic_pct', 'wood_pct', 'kitchen_pct', 'inert_pct',
                'other_pct', 'moisture_pct', 'fossil_share_of_C_pct', 'CO2_total_t_per_t_dry',
                'CO2_fossil_t_per_t_dry', 'residual_recomputed', 'n_facilities', 'n_moisture',
                'direct_incineration_FY2021_t'])
    for p in sorted(pref, key=lambda x: (pref[x]['region'], x)):
        d = pref[p]
        w.writerow([p, d['region']] + [round(d[k], 2) for k in KEYS]
                   + [round(d['moisture'], 2) if d['moisture'] is not None else '',
                      round(d['fossil_share'], 2), round(d['CO2_total'], 4), round(d['CO2_fossil'], 4),
                      d['residual_fixed'], d['n_fac'], d['n_moist'], int(direct[p])])

# ---------------------------------------------------------------- GSDs (same formula as R0)
old_pref = list(csv.DictReader(open(os.path.join(PRE, 'regional_source', 'prefecture_waste_carbon.csv'),
                                    encoding='utf-8-sig')))
G = {
    'formula': 'exp(sample standard deviation of ln x, ddof = 1)',
    'check_R0_fossil_share_47pref': gsd([float(p['fossil_share_of_C_pct']) for p in old_pref]),
    'check_R0_T_unweighted_regions': gsd([float(OLD_U[r]['CO2_total_t_per_t_dry']) for r in REGIONS]),
    'R1d_T_W2a_regions_not_used_by_R1d_MC': gsd([float(OLD_W2A[r]['CO2_total_t_per_t_dry']) for r in REGIONS]),
    'R1e_fossil_share_47pref': gsd([pref[p]['fossil_share'] for p in pref]),
    'R1e_T_W2a_regions': gsd([W2A[r]['CO2_total'] for r in REGIONS]),
    'R1e_fossil_share_W2a_regions': gsd([W2A[r]['fossil_share'] for r in REGIONS]),
    'R1e_CO2_fossil_W2a_regions': gsd([W2A[r]['CO2_fossil'] for r in REGIONS]),
    'MC_constants_R1d': {'throughput': 1.0129, 'fossil_share': 1.1062},
}
assert round(G['check_R0_fossil_share_47pref'], 4) == 1.1062 and round(G['check_R0_T_unweighted_regions'], 4) == 1.0129
json.dump(G, open(os.path.join(HERE, 'R1e_gsd.json'), 'w', encoding='utf-8'), indent=1)

# ---------------------------------------------------------------- preview through the chain's regional model
sys.path.insert(0, CODE)
os.environ.pop('R1C_SCENARIO', None)
import carbon_model as CB                                                      # noqa: E402
import regional_model as RMOD                                                  # noqa: E402
TABLES = {'U (submitted)': os.path.join(PRE, 'region_waste_carbon_unweighted.csv'),
          'R1d W2a': os.path.join(PRE, 'region_waste_carbon.csv'),
          'R1e W2a': os.path.join(HERE, 'region_waste_carbon_R1e_W2a.csv'),
          'R1e W1': os.path.join(HERE, 'region_waste_carbon_R1e_W1.csv')}
PREV = {}
for lab, path in TABLES.items():
    RMOD.REGION_FILES['_prev'] = path                     # absolute path: os.path.join keeps it
    regs = RMOD.load('_prev')
    res = {p: {n: CB.net(p, regs[n]['F'], regs[n]['T'], regs[n]['EF']) for n in regs} for p in CB.ORDER}
    better = {p: sum(1 for n in regs if res[p][n] < res['S0'][n]) for p in CB.ORDER if p != 'S0'}
    ben = {n: res['S0'][n] - res['G1'][n] for n in regs}
    PREV[lab] = dict(better=better, G1_benefit_min=min(ben.values()), G1_benefit_max=max(ben.values()),
                     G1_ratio=max(ben.values()) / min(ben.values()),
                     G1_minus_S0_range=[min(res['G1'][n] - res['S0'][n] for n in regs),
                                        max(res['G1'][n] - res['S0'][n] for n in regs)])
RMOD.REGION_FILES.pop('_prev')

# ---------------------------------------------------------------- comparison
rank = lambda d: {r: i + 1 for i, r in enumerate(sorted(d, key=lambda x: -d[x]))}
fsU = {r: float(OLD_U[r]['fossil_share_of_C_pct']) for r in REGIONS}
fsO = {r: float(OLD_W2A[r]['fossil_share_of_C_pct']) for r in REGIONS}
fsN = {r: W2A[r]['fossil_share'] for r in REGIONS}
fs1 = {r: W1[r]['fossil_share'] for r in REGIONS}
rU, rO, rN, r1 = rank(fsU), rank(fsO), rank(fsN), rank(fs1)
spear = lambda a, b: float(np.corrcoef([a[r] for r in REGIONS], [b[r] for r in REGIONS])[0, 1])
rows_cmp = []
for reg in REGIONS:
    o, n1 = OLD_W2A[reg], W2A[reg]
    rows_cmp.append([reg, round(fsU[reg], 2), round(fsO[reg], 2), round(fsN[reg], 2), round(fs1[reg], 2),
                     round(fsN[reg] - fsO[reg], 2), float(OLD_U[reg]['plastic_pct']), float(o['plastic_pct']),
                     round(n1['plas'], 2), round(W1[reg]['plas'], 2), float(o['CO2_total_t_per_t_dry']),
                     round(n1['CO2_total'], 4), float(o['CO2_fossil_t_per_t_dry']), round(n1['CO2_fossil'], 4),
                     rU[reg], rO[reg], rN[reg], r1[reg]])
with open(os.path.join(HERE, 'R1e_comparison.csv'), 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['region', 'fossil_U', 'fossil_R1d_W2a', 'fossil_R1e_W2a', 'fossil_R1e_W1', 'd_fossil_R1e_minus_R1d',
                'plastic_U', 'plastic_R1d_W2a', 'plastic_R1e_W2a', 'plastic_R1e_W1', 'CO2tot_R1d', 'CO2tot_R1e',
                'CO2fos_R1d', 'CO2fos_R1e', 'rank_U', 'rank_R1d', 'rank_R1e', 'rank_W1'])
    w.writerows(rows_cmp)

old_fs = {p['prefecture']: float(p['fossil_share_of_C_pct']) for p in old_pref}
dpref = sorted(((pref[p]['fossil_share'] - old_fs[p], p) for p in pref), key=lambda x: -abs(x[0]))

M = io.StringIO()
P = lambda s='': M.write(s + '\n')
P('# R1e 区域模块重建：新旧对照（R1d → R1e；这些表已安装到主链）')
P()
P('脚本 `build_R1e.py`，规则 `RULES_prefecture_composition.md`（R4 水分已按 2026-09-20 修订）。原始数据 '
  '`03_new_sources/MoEJ_FY2021/`。旧值（R1d）取自安装前的备份 `chain_r1/_snapshots/pre_R1e_2026-09-20/`。'
  '本脚本只写本文件夹；这些表随后已作为 R1e 安装到主链（见 `chain_r1/CHANGES_R1.md` R1e 一节），'
  '正式结果以主链重跑的 `P2C5_*` 为准。第 4 节的预览是用 `regional_model.load()` 和 `carbon_model.net()` 在内存中算出的。')
P()
P('## 1. 数据')
P()
P(f'- 设施：纳入 BD > 0 的 {len(facs)} / 1,060 座（R3 触发 {n_r3} 次）；47 县全部有纳入设施。'
  f'有效水分（BM > 0）的县：{47 - len(missing_moist)} / 47' + (f'，缺失：{"、".join(missing_moist)}' if missing_moist else '') + '。')
P(f'- W1 使用 BD > 0 且年間処理量 > 0 的 {n_w1} 座设施，处理量合计 {thr_w1:,.0f} t/年度，'
  f'为全国直接焼却量 {NAT:,.0f} t 的 {100 * thr_w1 / NAT:.1f} %。')
P(f'- 县级化石碳占比变化最大的 5 县（R1e − 投稿版县值，pp）：'
  + '；'.join(f'{p} {d:+.2f}' for d, p in dpref[:5]) + '。')
P()
P('## 2. 十区域：化石碳占比（%）与排序（1 = 最高）')
P()
P('| 区域 | U（投稿版） | R1d W2a | **R1e W2a** | R1e W1 | Δ R1e − R1d | 塑料 R1d → **R1e** | CO₂ fossil t/t R1d → **R1e** | 排序 U / R1d / **R1e** / W1 |')
P('|---|---:|---:|---:|---:|---:|---|---|---|')
for c in rows_cmp:
    P(f'| {c[0]} | {c[1]:.2f} | {c[2]:.2f} | **{c[3]:.2f}** | {c[4]:.2f} | {c[5]:+.2f} | {c[7]:.2f} → **{c[8]:.2f}** | '
      f'{c[12]:.4f} → **{c[13]:.4f}** | {c[14]} / {c[15]} / **{c[16]}** / {c[17]} |')
P()
P(f'区域化石碳占比范围：U {min(fsU.values()):.2f}–{max(fsU.values()):.2f}；R1d {min(fsO.values()):.2f}–{max(fsO.values()):.2f}；'
  f'**R1e {min(fsN.values()):.2f}–{max(fsN.values()):.2f}**；W1 {min(fs1.values()):.2f}–{max(fs1.values()):.2f}。'
  f'排序相关（Spearman，基于名次）：R1e 与 R1d {spear(rN, rO):.3f}，与 U {spear(rN, rU):.3f}，与 W1 {spear(rN, r1):.3f}。')
P()
P('## 3. MC 宽度（exp(ln 值样本标准差)，ddof = 1）')
P()
P('| 量 | 投稿版 / R1d MC 常数 | 公式复现旧值 | **R1e 新值（完整精度）** |')
P('|---|---:|---:|---:|')
P(f'| fossil_share（47 县） | 1.1062 | {G["check_R0_fossil_share_47pref"]:.6f} | **{G["R1e_fossil_share_47pref"]:.6f}** |')
P(f'| throughput T（10 区域 CO₂ total） | 1.0129（取自不加权表） | {G["check_R0_T_unweighted_regions"]:.6f} | '
  f'**{G["R1e_T_W2a_regions"]:.6f}**（R1d W2a 表本应为 {G["R1d_T_W2a_regions_not_used_by_R1d_MC"]:.6f}） |')
P(f'| 参考：区域化石碳占比 / 区域 CO₂ fossil | 1.079 / 1.093（注释） | — | {G["R1e_fossil_share_W2a_regions"]:.4f} / {G["R1e_CO2_fossil_W2a_regions"]:.4f} |')
P()
P('按已商定的做法：无论显示精度上是否相同，下一步都用完整精度更新 MC 常数，并重跑 50,000 次。')
P()
P('## 4. 预览：2025 年区域结果（主链区域模型，内存计算，口径 A，未写回）')
P()
P('| 指标 | ' + ' | '.join(TABLES) + ' |')
P('|---|' + '---:|' * len(TABLES))
for p in [q for q in CB.ORDER if q != 'S0']:
    P(f'| {p} 优于 S0 的区域数 /10 | ' + ' | '.join(str(PREV[t]['better'][p]) for t in TABLES) + ' |')
P('| G1 区域收益 S0 − G1（t/y）最小–最大 | ' + ' | '.join(f'{PREV[t]["G1_benefit_min"]:,.0f}–{PREV[t]["G1_benefit_max"]:,.0f}' for t in TABLES) + ' |')
P('| G1 收益比（最大/最小） | ' + ' | '.join(f'{PREV[t]["G1_ratio"]:.2f}' for t in TABLES) + ' |')
P()
_f = json.load(open(os.path.join(ANA, 'chain_r1', 'P2C4_key_numbers.json'), encoding='utf-8'))
assert abs(PREV['R1d W2a']['G1_benefit_min'] / 1000 - _f['G1_benefit_min'][-1]) < 1e-3 and abs(PREV['R1d W2a']['G1_ratio'] - _f['G1_benefit_ratio'][-1]) < 1e-9
P(f'核对：R1d W2a 一列与 R1d 冻结值一致（G1 收益 {_f["G1_benefit_min"][-1]:.3f}–{_f["G1_benefit_max"][-1]:.3f} kt/y，比值 '
  f'{_f["G1_benefit_ratio"][-1]:.2f}），说明预览与正式链的定义相同。2050 年区域计数、口径 B 和 MC 要等正式重跑。')
open(os.path.join(HERE, 'R1e_comparison.md'), 'w', encoding='utf-8').write(M.getvalue())
print(M.getvalue())
