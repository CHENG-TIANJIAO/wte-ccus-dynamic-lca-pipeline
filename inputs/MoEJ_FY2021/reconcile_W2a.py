# -*- coding: utf-8 -*-
"""
Rebuild the tonnage weights of the regional layer (W2a) from the raw MoEJ FY2021 files in
this folder, then check them against the table the calculation chain actually uses.

Raw files (downloaded 2026-09-20 00:09 JST, see SOURCE_RECORD.md for URLs and SHA-256):
  r3_shori_total_02.xlsx     全体集計結果（ごみ処理状況）, sheet ごみ処理概要, P3 直接焼却量, unit P6 （ｔ）,
                             rows 7-53 = 47 prefectures, row 54 = 全国
  r3_seibi_facility_01.xlsx  焼却施設, sheet 焼却, G2 年間処理量 (t/年度), one row per facility from row 7

Chain inputs compared against (read-only):
  chain_r1/calculation_code/regional_source/prefecture_waste_carbon.csv  47 prefecture profiles
  chain_r1/calculation_code/region_waste_carbon.csv                     W2a table used by the chain

Checks
  1. 47 prefecture rows, names match the chain's prefecture list; P54 全国 = SUM(P7:P53) = 31,491,040 t.
  2. Facility table: 1,060 facility rows, 47 prefectures.
  3. Region weights: sum of 直接焼却量 per region and share of the national total equal the chain's
     direct_incineration_FY2021_t / share_of_national_direct_incineration_pct.
  4. Region values: 直接焼却量-weighted means of the prefecture profiles reproduce the chain's
     plastic_pct, moisture_pct, fossil_share_of_C_pct, CO2_total_t_per_t_dry, CO2_fossil_t_per_t_dry
     (to the rounding of the shipped prefecture CSV).
Writes W2a_weights_rebuilt.csv (prefecture -> region -> weight) and prints the reconciliation.
Run:  python reconcile_W2a.py
"""
import csv
import hashlib
import os
import sys

import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
HERE = os.path.dirname(os.path.abspath(__file__))
CHAIN = os.path.normpath(os.path.join(HERE, '..', '..', '02_analyses', 'chain_r1', 'calculation_code'))
TOT = os.path.join(HERE, 'r3_shori_total_02.xlsx')
FAC = os.path.join(HERE, 'r3_seibi_facility_01.xlsx')
SHA = {'r3_shori_total_02.xlsx': 'b12aefdffcf1ec30c8b276281f28ac218e271eb577cb3715dc53a9af49d2aac5',
       'r3_seibi_facility_01.xlsx': '77a95fe0dbe53dc13ef579c9b49d9a9fcc5cbe523c7f49b37a5b28372abd0fe8'}
for f, h in SHA.items():
    assert hashlib.sha256(open(os.path.join(HERE, f), 'rb').read()).hexdigest() == h, f
print('0. SHA-256 of both raw files = SOURCE_RECORD.md')

# ---------------------------------------------------------------- 1. prefecture 直接焼却量
ws = openpyxl.load_workbook(TOT, read_only=True, data_only=False)['ごみ処理概要']
assert ws['P3'].value == '直接焼却量' and ws['P6'].value == '（ｔ）', (ws['P3'].value, ws['P6'].value)
direct = {}
for r in range(7, 54):
    direct[ws.cell(r, 1).value] = float(ws.cell(r, 16).value)
nat_label, nat_formula = ws.cell(54, 1).value, ws.cell(54, 16).value
wv = openpyxl.load_workbook(TOT, read_only=True, data_only=True)['ごみ処理概要']
nat_value = float(wv.cell(54, 16).value)
assert len(direct) == 47
assert abs(sum(direct.values()) - 31491040) < 0.5 and abs(nat_value - 31491040) < 0.5
print(f'1. 直接焼却量: 47 prefectures (rows 7-53), sum {sum(direct.values()):,.0f} t; row 54 "{nat_label}" '
      f'= {nat_value:,.0f} t (formula {nat_formula})')

prof = list(csv.DictReader(open(os.path.join(CHAIN, 'regional_source', 'prefecture_waste_carbon.csv'),
                                encoding='utf-8-sig')))
assert sorted(p['prefecture'] for p in prof) == sorted(direct), \
    set(direct) ^ {p['prefecture'] for p in prof}
print('   prefecture names identical to the chain profiles (47)')

# ---------------------------------------------------------------- 2. facility table
wf = openpyxl.load_workbook(FAC, read_only=True, data_only=True)['焼却']
assert wf['G2'].value == '年間処理量' and wf['G6'].value == '（t/年度）', (wf['G2'].value, wf['G6'].value)
rows = [r for r in wf.iter_rows(min_row=7, values_only=True) if r[0]]
prefs_fac = {r[0] for r in rows}
thr = sum(float(r[6]) for r in rows if isinstance(r[6], (int, float)))
assert len(rows) == 1060 and len(prefs_fac) == 47, (len(rows), len(prefs_fac))
print(f'2. facility table 焼却: {len(rows):,} facility rows (row 7 on), {len(prefs_fac)} prefectures, '
      f'sum 年間処理量 {thr:,.0f} t/年度 ({100 * thr / nat_value:.1f} % of 直接焼却量)')

# ---------------------------------------------------------------- 3./4. regions
chain = {r['region']: r for r in csv.DictReader(open(os.path.join(CHAIN, 'region_waste_carbon.csv'),
                                                      encoding='utf-8-sig'))}
COLS = [('plastic_pct', 2), ('moisture_pct', 2), ('fossil_share_of_C_pct', 2),
        ('CO2_total_t_per_t_dry', 4), ('CO2_fossil_t_per_t_dry', 4)]
out, worst = [], {c: 0.0 for c, _ in COLS}
print('3./4. region    n  直接焼却量 t (rebuilt = chain)   share %   ' + '  '.join(c.split('_')[0] for c, _ in COLS))
for reg in sorted(chain):
    ms = [p for p in prof if p['region'] == reg]
    w = [direct[p['prefecture']] for p in ms]
    W = sum(w)
    ch = chain[reg]
    assert len(ms) == int(ch['n_prefectures'])
    assert abs(W - float(ch['direct_incineration_FY2021_t'])) < 0.5, (reg, W, ch['direct_incineration_FY2021_t'])
    share = 100 * W / nat_value
    assert abs(round(share, 2) - float(ch['share_of_national_direct_incineration_pct'])) < 1e-9, (reg, share)
    vals = []
    for c, nd in COLS:
        v = sum(float(p[c]) * wi for p, wi in zip(ms, w)) / W
        d = abs(v - float(ch[c]))
        worst[c] = max(worst[c], d)
        vals.append(f'{v:.{nd}f}/{float(ch[c]):.{nd}f}')
    print(f'   {reg:<9}{len(ms):>3}  {W:>14,.0f}                  {share:6.2f}   ' + '  '.join(vals))
    for p, wi in zip(ms, w):
        out.append([p['prefecture'], reg, int(wi), round(100 * wi / W, 4), round(100 * wi / nat_value, 4)])
# bound: the prefecture inputs are rounded (+-0.5 unit of the last decimal, so is their weighted mean)
# and the chain's region values are rounded again (+-0.5 unit) -> |difference| <= 1 unit of the last decimal
TOL = {'plastic_pct': 0.01, 'moisture_pct': 0.01, 'fossil_share_of_C_pct': 0.01,
       'CO2_total_t_per_t_dry': 1e-4, 'CO2_fossil_t_per_t_dry': 1e-4}
print('   max |rebuilt - chain|: ' + ', '.join(f'{c} {worst[c]:.2g}' for c, _ in COLS))
print('   (tolerance = one unit of the last decimal: rounded prefecture inputs + rounded region outputs; the '
      'exact, unrounded rebuild of the prefecture profiles and regions is '
      '02_analyses/R1e_regional/build_R1e.py, whose outputs are the installed chain tables byte for byte)')
assert all(worst[c] <= TOL[c] for c in worst), worst

with open(os.path.join(HERE, 'W2a_weights_rebuilt.csv'), 'w', newline='', encoding='utf-8') as f:
    wr = csv.writer(f)
    wr.writerow(['prefecture', 'region', 'direct_incineration_FY2021_t', 'weight_within_region_pct',
                 'share_of_national_pct'])
    wr.writerows(sorted(out, key=lambda x: (x[1], x[0])))
print(f'wrote W2a_weights_rebuilt.csv ({len(out)} prefectures). ALL CHECKS PASSED')
