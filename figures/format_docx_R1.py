# -*- coding: utf-8 -*-
"""Post-build typography of the R1 manuscript and SI (format pass, 2026-09-20).

Chemical formulae get subscript digits (CO2 -> CO₂ as a subscript run, CH4, CH3OH, H2, H2O, N2O, NH3, CaCO3,
MgCO3, Mg2SiO4, FeCl3, SO2) and unit exponents become superscripts with a true minus (y-1 -> y⁻¹ as a
superscript run, likewise t, d, km, kWh, MWh, kg, g, m, MJ, GJ).  Each match is split into its own run that
copies the formatting of the original run, so bold, italics and size are kept.  Formulae inside symbols
(B_H2, m_H2) and codes (U2a, G2, M1-M4) are not touched: a formula must stand alone, and an exponent must follow
a unit that stands alone.  Runs holding drawings are skipped.  Returns the counts."""
import copy
import re
import sys

import docx
from docx.text.run import Run

CHEM = ['NH4Cl', 'CH3OH', 'Mg2SiO4', 'CaCO3', 'MgCO3', 'FeCl3', 'CO2', 'CH4', 'H2O', 'N2O', 'NH3', 'NH4', 'SO2', 'H2']
CHEM_RE = re.compile(r'(?<![A-Za-z0-9_])(' + '|'.join(CHEM) + r')(?![A-Za-z0-9_])')
UNITS = ['kWh', 'MWh', 'GWh', 'km', 'kg', 'MJ', 'GJ', 'y', 't', 'd', 'g', 'm']
EXP_RE = re.compile(r'(?<![A-Za-z0-9_])(' + '|'.join(UNITS) + r')[-−]([123])(?!\d)(?![.,]\d)')   # 'y-1,' yes; 'y-1.5', 'y-1,000' no


EQ_LINE = re.compile(r'\((?:S?\d+[a-d]?)\)\s*$')


def segments(text, units=True):
    """Split text into (piece, vert) with vert in {None, 'sub', 'sup'}.

    `units` is False inside an equation, where a unit exponent cannot occur but the same characters can:
    F(t-1) is the value one year earlier, not F per tonne."""
    marks = []
    for m in CHEM_RE.finditer(text):
        f = m.group(1)
        for k in re.finditer(r'\d+', f):
            marks.append((m.start(1) + k.start(), m.start(1) + k.end(), 'sub', None))
    for m in (EXP_RE.finditer(text) if units else ()):
        marks.append((m.end(1), m.end(2), 'sup', '−' + m.group(2)))
    if not marks:
        return None
    marks.sort()
    out, pos = [], 0
    for a, b, v, rep in marks:
        if a > pos:
            out.append((text[pos:a], None))
        out.append((rep if rep else text[a:b], v))
        pos = b
    if pos < len(text):
        out.append((text[pos:], None))
    return out


def fix_paragraph(p, counts, units=True):
    for r in list(p.runs):
        if r._r.xpath('.//w:drawing') or not r.text:
            continue
        seg = segments(r.text, units)
        if not seg:
            continue
        base = copy.deepcopy(r._r)
        anchor = r._r
        first = True
        for piece, v in seg:
            if first:
                el = anchor
                first = False
            else:
                el = copy.deepcopy(base)
                anchor.addnext(el)
                anchor = el
            run = Run(el, p)
            run.text = piece
            if v == 'sub':                           # sub- and superscript share w:vertAlign: set one only
                run.font.subscript = True
            elif v == 'sup':
                run.font.superscript = True
            if v:
                counts[v] += 1


def walk(container, counts):
    for p in container.paragraphs:
        fix_paragraph(p, counts, units=not EQ_LINE.search(p.text.strip()))
    for t in container.tables:
        for row in t.rows:
            for cell in row.cells:
                walk(cell, counts)


def format_docx(path):
    d = docx.Document(path)
    counts = {'sub': 0, 'sup': 0}
    walk(d, counts)
    for s in d.sections:
        for part in (s.header, s.footer):
            walk(part, counts)
    d.save(path)
    return counts


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    for pth in sys.argv[1:]:
        print(pth, format_docx(pth))
